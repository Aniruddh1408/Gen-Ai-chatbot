"""Extract plain text from uploaded documents.

The deployed LLM is text-only, so every uploaded file must be turned into text
before the assistant can answer about it. This module reads the supported
document types (PDF, Word, text/markdown, CSV, Excel) and returns clean text.

Design rules:
- Never raise to the caller. Parsing failures return an ``ExtractResult`` with a
  safe, user-facing ``error`` message; technical detail is logged.
- Heavy/optional parsers (pypdf, python-docx, openpyxl) are imported lazily so
  the module loads even if a parser is missing; a missing parser yields a
  friendly error rather than a crash.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from app.core.config import UPLOAD_ALLOWED_EXTENSIONS, UPLOAD_MAX_BYTES
from app.core.logger import get_logger

logger = get_logger("file_extract")

# Excel cell/row caps so a huge spreadsheet can't blow up memory or the prompt.
_XLSX_MAX_ROWS = 2000
_CSV_MAX_ROWS = 5000


@dataclass
class ExtractResult:
    """Outcome of a text-extraction attempt."""
    text: str = ""
    n_units: int = 0          # pages (pdf), sheets (xlsx), rows (csv), or paragraphs
    kind: str = ""            # normalized extension, e.g. "pdf"
    error: str | None = None  # safe user-facing message, or None on success

    @property
    def ok(self) -> bool:
        return self.error is None and bool(self.text.strip())


def file_extension(filename: str) -> str:
    """Return the lowercase extension without a leading dot."""
    name = (filename or "").strip().lower()
    _, _, ext = name.rpartition(".")
    return ext if ext and ext != name else ""


def _decode_bytes(raw: bytes) -> str:
    """Decode bytes as UTF-8, falling back gracefully for other encodings."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, AttributeError):
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_txt(raw: bytes) -> ExtractResult:
    text = _decode_bytes(raw).strip()
    return ExtractResult(text=text, n_units=text.count("\n") + 1 if text else 0, kind="txt")


def _extract_pdf(raw: bytes) -> ExtractResult:
    try:
        from pypdf import PdfReader
    except Exception as exc:  # pragma: no cover - depends on optional dep
        logger.error(f"pypdf unavailable: {exc}")
        return ExtractResult(kind="pdf", error="PDF support is not available on the server right now.")

    try:
        reader = PdfReader(io.BytesIO(raw))
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                pages.append("")
        text = "\n\n".join(p.strip() for p in pages if p.strip()).strip()
        return ExtractResult(text=text, n_units=len(reader.pages), kind="pdf")
    except Exception as exc:
        logger.error(f"PDF extraction failed: {exc}")
        return ExtractResult(kind="pdf", error="I couldn't read that PDF. It may be corrupted or password-protected.")


def _extract_docx(raw: bytes) -> ExtractResult:
    try:
        import docx  # python-docx
    except Exception as exc:  # pragma: no cover - depends on optional dep
        logger.error(f"python-docx unavailable: {exc}")
        return ExtractResult(kind="docx", error="Word document support is not available on the server right now.")

    try:
        document = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
        # Include table cell text, which python-docx keeps outside paragraphs.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        return ExtractResult(text=text, n_units=len(parts), kind="docx")
    except Exception as exc:
        logger.error(f"DOCX extraction failed: {exc}")
        return ExtractResult(kind="docx", error="I couldn't read that Word document. It may be corrupted.")


def _extract_csv(raw: bytes) -> ExtractResult:
    try:
        decoded = _decode_bytes(raw)
        reader = csv.reader(io.StringIO(decoded))
        lines = []
        for index, row in enumerate(reader):
            if index >= _CSV_MAX_ROWS:
                lines.append(f"... (truncated at {_CSV_MAX_ROWS} rows)")
                break
            lines.append(", ".join(cell.strip() for cell in row))
        text = "\n".join(line for line in lines if line.strip()).strip()
        return ExtractResult(text=text, n_units=len(lines), kind="csv")
    except Exception as exc:
        logger.error(f"CSV extraction failed: {exc}")
        return ExtractResult(kind="csv", error="I couldn't read that CSV file.")


def _extract_xlsx(raw: bytes) -> ExtractResult:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - depends on optional dep
        logger.error(f"openpyxl unavailable: {exc}")
        return ExtractResult(kind="xlsx", error="Excel support is not available on the server right now.")

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet_count = len(workbook.sheetnames)
        blocks = []
        row_budget = _XLSX_MAX_ROWS
        for sheet in workbook.worksheets:
            blocks.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                if row_budget <= 0:
                    blocks.append("... (truncated)")
                    break
                cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if cells:
                    blocks.append(", ".join(cells))
                    row_budget -= 1
        workbook.close()
        text = "\n".join(blocks).strip()
        return ExtractResult(text=text, n_units=sheet_count, kind="xlsx")
    except Exception as exc:
        logger.error(f"XLSX extraction failed: {exc}")
        return ExtractResult(kind="xlsx", error="I couldn't read that Excel file. It may be corrupted.")


_EXTRACTORS = {
    "txt": _extract_txt,
    "md": _extract_txt,
    "pdf": _extract_pdf,
    "docx": _extract_docx,
    "csv": _extract_csv,
    "xlsx": _extract_xlsx,
}


def extract_text(filename: str, raw: bytes) -> ExtractResult:
    """Extract text from an uploaded file's raw bytes.

    Validates extension and size, dispatches to the right parser, and always
    returns an ``ExtractResult`` (never raises).
    """
    ext = file_extension(filename)
    if not ext or ext not in UPLOAD_ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(UPLOAD_ALLOWED_EXTENSIONS))
        return ExtractResult(kind=ext, error=f"Unsupported file type. Allowed types: {allowed}.")

    if not raw:
        return ExtractResult(kind=ext, error="The file appears to be empty.")

    if len(raw) > UPLOAD_MAX_BYTES:
        max_mb = UPLOAD_MAX_BYTES / (1024 * 1024)
        return ExtractResult(kind=ext, error=f"File is too large. Maximum size is {max_mb:.0f} MB.")

    extractor = _EXTRACTORS.get(ext)
    if extractor is None:  # pragma: no cover - guarded by allowlist above
        return ExtractResult(kind=ext, error="Unsupported file type.")

    result = extractor(raw)
    if result.error is None and not result.text.strip():
        result.error = "I couldn't find any readable text in that file. (Scanned images need OCR, which isn't enabled.)"
    return result
